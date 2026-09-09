#!/usr/bin/env python3
"""Export the money model from Goods-financial-plan.xlsx to Notion-flavoured Markdown.

One direction only: workbook -> page. Recalculates a copy with LibreOffice first, because the
master is saved without cached values. Usage: python3 export-notion-model.py > model.md
"""
import openpyxl, subprocess, tempfile, shutil, os, hashlib, datetime, glob, sys
HERE=os.path.dirname(os.path.abspath(__file__)); SRC=os.path.join(HERE,"Goods-financial-plan.xlsx")
sha=hashlib.sha256(open(SRC,"rb").read()).hexdigest()[:8]
tmp=tempfile.mkdtemp(); shutil.copy(SRC,os.path.join(tmp,"in.xlsx"))
subprocess.run(["/opt/homebrew/bin/soffice","--headless","--calc","--convert-to",'xlsx:Calc MS Excel 2007 XML',"--outdir",os.path.join(tmp,"out"),os.path.join(tmp,"in.xlsx")],capture_output=True,timeout=240)
wb=openpyxl.load_workbook(glob.glob(os.path.join(tmp,"out","*.xlsx"))[0],data_only=True)
F,P,Q,C,A,S,As,R=(wb[n] for n in ["Funding","Phase and raise","QBE","Cases","Alignment","Sites","Assumptions","Responsibilities"])
def m(v):
    if v is None or v=="": return ""
    try: return f"${float(v):,.0f}"
    except: return str(v)
def table(headers,rows):
    out=['<table header-row="true">','<tr>'+"".join(f"<td>{h}</td>" for h in headers)+'</tr>']
    for r in rows: out.append('<tr>'+"".join(f"<td>{'' if c is None else c}</td>" for c in r)+'</tr>')
    out.append('</table>'); return "\n".join(out)
today=datetime.date.today().strftime("%-d %B %Y")
L=[]
L.append(f"Generated from Goods-financial-plan.xlsx on {today} (SHA {sha}). Do not edit numbers here; change the workbook and push again. AUD; buyer receipts incl. GST, model costs ex GST.")
L.append("## The model in five layers")
L.append("Goods on Country transfers manufacturing and economic ownership into communities, using beds as the first proven product. The goal is not to produce beds forever; it is to produce community-owned enterprises. 1 Community demand: map need and buyers. 2 Catalytic enterprise capital: philanthropy pays Goods $750 a bed for Witta-made first trading stock. 3 Institutional demand: buyers with existing health and housing budgets buy beds from the community enterprise; the money stays with the community, never Goods revenue. 4 Manufacturing transfer: QBE's plants; Alice Springs through DEWR. 5 Shared services and governance. Each dollar has a job. Success is enterprises operating and revenue kept locally, not beds made.")
L.append("## Every way a bed gets paid for (Phase and raise rows 112 to 128)")
rows=[]
for r in range(113,122):
    if P.cell(r,1).value: rows.append([str(P.cell(r,4).value),P.cell(r,1).value,m(P.cell(r,2).value),str(round(P.cell(r,3).value)) if isinstance(P.cell(r,3).value,(int,float)) else "",P.cell(r,5).value])
L.append(table(["Tier","Source","$","Beds at $750","Status"],rows))
L.append(table(["Beds funded","Beds","Meaning"],[["Tier 1, invited",str(round(P["B123"].value)),"The floor"],["Tiers 1 + 2, plus asks to make",str(round(P["B124"].value)),"This year's plan; "+str(P["B126"].value)+" communities at 100 each"],["Tiers 1 + 2 + 3, plus repayable",str(round(P["B125"].value)),"Only where the loan has a repayment source. Witta's ceiling is "+str(round(P["B127"].value))+"."]]))
L.append("Ramp up or slow down from this ladder, never from the cost side. Tiers 1 and 2 amounts for Snow and Dusseldorp are ours, not theirs, until an ask exists.")
L.append("## Can Witta make them, and when do the plants need to be live? (rows 131 to 144)")
L.append(table(["Period","Witta, own panels","Witta, Defy legs","Demand on Witta","Reads"],[[P.cell(r,1).value,str(round(P.cell(r,2).value)),str(round(P.cell(r,3).value)),str(P.cell(r,4).value),P.cell(r,5).value] for r in range(133,136)]))
L.append(table(["Plant","First beds","Year one","Steady, a year","Basis"],[[P.cell(r,1).value,P.cell(r,2).value,str(P.cell(r,3).value),str(P.cell(r,4).value),P.cell(r,5).value] for r in range(138,143)]))
L.append(str(P["A144"].value))
L.append("## The shred loop (rows 146 to 156)")
L.append(table(["","kg","Meaning"],[[P.cell(r,1).value,str(P.cell(r,2).value) if P.cell(r,2).value is not None else "not yet measured",P.cell(r,4).value] for r in range(147,157)]))
BC=wb["Bed calculator"]
def d(v):
    try: return v.strftime("%b %Y")
    except: return str(v) if v is not None else ""
def n(v):
    try: return f"{float(v):,.0f}"
    except: return str(v) if v is not None else ""
L.append("## Bed calculator: what we have, what to order, when a plant is needed (Bed calculator sheet)")
L.append("Yellow cells on the sheet are inputs. Count of 9 Sep 2026; bucket and skeleton weights are estimates until weighed.")
L.append(table(["Stock at Witta","Count"],[[BC.cell(r,1).value,n(BC.cell(r,2).value)] for r in range(5,21)]))
L.append(table(["What that makes","Beds"],[[BC.cell(r,1).value,n(BC.cell(r,2).value) if not isinstance(BC.cell(r,2).value,str) else BC.cell(r,2).value] for r in range(23,31)]))
L.append(table(["Month","Demand","Witta own","Plants","Capacity own + plants","Cumulative demand","Cumulative capacity","Short","Shred needed, kg","Bags to order"],[[d(BC.cell(r,1).value),n(BC.cell(r,2).value),n(BC.cell(r,3).value),n(BC.cell(r,5).value+BC.cell(r,6).value+BC.cell(r,7).value),n(BC.cell(r,8).value),n(BC.cell(r,9).value),n(BC.cell(r,10).value),n(BC.cell(r,11).value),n(BC.cell(r,13).value),n(BC.cell(r,16).value)] for r in range(55,67)]))
L.append(table(["Reads","Value"],[[BC.cell(r,1).value, d(BC.cell(r,2).value) if r in (69,72) else n(BC.cell(r,2).value)] for r in (69,70,71,72,73,74)]))
L.append(table(["What to order now, October to December","$ or count"],[[BC.cell(r,1).value,n(BC.cell(r,2).value)] for r in range(77,89)]))
L.append(table(["Money across every area for the year","$"],[[BC.cell(r,1).value,n(BC.cell(r,2).value)] for r in range(91,105)]))
L.append(table(["First principles","Value"],[[BC.cell(r,1).value, d(BC.cell(r,2).value) if r==114 else (BC.cell(r,2).value if isinstance(BC.cell(r,2).value,str) else n(BC.cell(r,2).value))] for r in range(107,115)]))
PJ=wb["Plan to July 2027"]
L.append("## Plan to July 2027: every dollar in and out by month (Plan to July 2027 sheet)")
L.append("Buyers who pay a community enterprise are not in this table by design; that money never reaches Goods. ALIVE's 100 beds are already paid for, so they leave stock in November with no cash in and $474 a bed already Goods' to keep.")
L.append(table(["Money landing","Month","Amount","Included","Kind","Basis"],[[PJ.cell(r,1).value,d(PJ.cell(r,2).value),n(PJ.cell(r,3).value),str(PJ.cell(r,4).value),PJ.cell(r,5).value,PJ.cell(r,6).value] for r in range(22,32)]))
L.append(table(["Month","Beds made","Grants and catalytic in","Buyers paying Goods","Make","Freight","Running","Facilitation","Plants","Net","Closing cash","Contribution earned"],[[d(PJ.cell(r,1).value)]+[n(PJ.cell(r,j).value) for j in (2,3,4,7,8,9,10,11,13,14,15)] for r in range(40,51)]))
L.append(table(["Read","Value"],[[PJ.cell(r,1).value, d(PJ.cell(r,2).value) if r==52 else (PJ.cell(r,2).value if isinstance(PJ.cell(r,2).value,str) else n(PJ.cell(r,2).value))] for r in (51,52,53,57,58,59,60,62,63,64,65,66,69,70,71,72,73,74,75)]))
L.append("## The asks, and the four outcomes each buys")
L.append(table(["Funder","Ask","What it buys","Beds","Labour hours","Communities at 100","$ kept locally when sold","kg diverted","Beds off the floor","Status"],[[PJ.cell(r,1).value,n(PJ.cell(r,2).value),PJ.cell(r,3).value]+[str(PJ.cell(r,j).value) if isinstance(PJ.cell(r,j).value,str) else n(PJ.cell(r,j).value) for j in (4,5,6,7,8,9)]+[PJ.cell(r,10).value] for r in range(84,93)]))
L.append(str(PJ.cell(93,1).value))
L.append("## Five and ten years, from the same three numbers")
L.append(table(["Year","Plants","Beds","Revenue to Goods","Stays in Goods","Running the business","Result","Kept in communities","kg diverted","Labour hours","Communities"],[[PJ.cell(r,1).value]+[n(PJ.cell(r,j).value) for j in range(2,12)] for r in range(103,113)]))
L.append(str(PJ.cell(114,1).value))
L.append("## 1. Who pays for what")
L.append(table(["Holder","Money","Buys","Status"],[
 ["QBE Foundation (Catalysing Impact)",m(As["B35"].value)+" (smaller "+m(As["B36"].value)+")","The two community plants outside Alice Springs, working choice Palm Island and Maningrida (Mount Isa, Tennant Creek, Ceduna in the mix)","Application due Fri 25 Sep 12pm; review 7 Oct; conditional outcome 23 Oct; pre-conditions 13 Nov"],
 ["Tim Fairfax Family Foundation","$100,000 a year for three years","The organisation: shared Goods costs (general operating support, Resilience stream)","Invited 31 Aug; apply by 9 Oct; board late Nov"],
 ["Brian M. Davis Charitable Foundation","up to $100,000, 12 months","Facilitation in four bed communities ($40,000) and beds ($60,000)","Invited 1 Sep; apply by 25 Sep; board 19 Nov; pays only once all project funding is confirmed"],
 ["Oonchiumpa (DEWR REAL Innovation Fund)","$1,695,000 over four years; about $150,000 to Goods","The Alice Springs plant, its project manager and wages. Outside this raise.","Offer 12 Aug; Grant Agreement pending"],
 ["Oonchiumpa (NIAA Local Investments)","up to $150,000","Alice Springs plant equipment, paid to Goods. Outside this raise.","Invitation in train; NIAA confirmed equipment eligible"],
 ["Sefa (Backing the Bold)","$200,000 target, ours not theirs","Working capital or plant, repayable","No EOI yet; Joel Bird will review one"],
 ["SEDI First Nations / general SEDI","$50,000 to $120,000","Capability services only (finance, evaluation, legal). Cannot buy plant or beds.","Rolling EOI to early 2027; applicant to choose"],
 ["Institutional buyers (layer 3)","see section 3","Beds from the community enterprise; the money stays with the community. ALIVE's 100 (paid July) are a separate order outside the 400.","Centrecorp 130 on quote; NT Health and Miwatj, Shepherdson College, Ceduna Aboriginal Corporation in conversation"],
 ["The organisation and factory year",m(P["B92"].value)+" a year","Founders, travel, accounting, marketing, Witta rent and upkeep. No coordinator, no trainer, no head office (Ben, 9 Sep). Carried by what stays in every bed ($474 before freight), TFFF operating support and Goods' other income. FY26 Goods-tagged receipts: "+m(P["B94"].value)+".","Inside the $750, not a separate raise"]]))
L.append("## The price model: every bed carries the business")
L.append(table(["","Amount"],[["Price per bed",m(P["B87"].value)],["Make it at Witta (canon $276: plastic, diesel, factory labour, steel, canvas, hardware). Price is at the factory door; freight is charged on top at cost, per community",m(P["B88"].value)],["Stays in Goods per bed, for rent, founders and everything else (freight is on top, paid with the bed)",m(P["B89"].value)],["400 beds: revenue to Goods",m(P["B90"].value)],["400 beds: stays in Goods",m(P["B91"].value)],["Running the business a year: founders, travel, accounting, marketing, Witta rent and upkeep",m(P["B92"].value)],["Beds a year to carry the organisation on bed sales alone",str(P["B93"].value)],["Goods-tagged receipts FY26 (beds, washers, grants)",m(P["B94"].value)]]))
L.append("Who pays Goods for the 400 first-stock beds (catalytic capital, layer 2): Tim Fairfax year one 133 ($100,000, invited), Brian M. Davis 80 ($60,000, invited; its other $40,000 pays facilitation in four communities), and 187 beds ($140,000) still to find from philanthropy or a smaller first year of 213. ALIVE's 100 are a separate paid order outside the 400, delivered around Central Australia under ALIVE's own grant. Institutional buyers (layer 3: Centrecorp 130 on quote, NT Health and Miwatj, Shepherdson College, Ceduna Aboriginal Corporation) buy beds from the community enterprise and that money stays with the community; it is never Goods revenue. QBE builds the two plants ($300,000), layer 4, outside the bed price.")
L.append("### Full cost of a bed by volume")
L.append(table(["Beds a year","Make it","Rent, founders, facilitation and the rest, per bed, if beds were the only thing Goods sold","Full cost per bed","Against $750"],[[f"{int(P.cell(rr,2).value):,}",m(P.cell(rr,3).value),m(P.cell(rr,4).value),m(P.cell(rr,5).value),m(P.cell(rr,6).value)] for rr in range(105,109)]))
L.append("If beds were the only thing Goods sold, about 630 a year would carry the whole business. They are not the only thing: washing machines, workshops and grants carry their share too. Read this table as a ceiling, not a target.")
L.append("## 2. Incoming: every live source")
rows=[]
for r in range(6,19):
    if F.cell(r,2).value: rows.append([F.cell(r,1).value,F.cell(r,2).value,F.cell(r,3).value,F.cell(r,4).value,F.cell(r,5).value,m(F.cell(r,6).value),"Yes" if F.cell(r,10).value==1 else "No",F.cell(r,12).value])
L.append(table(["ID","Funder","Instrument","Recipient","Stage","Year 1 amount","In Full case?","Condition / next evidence"],rows))
L.append(f"Secured today: {m(F['F19'].value)}. Committed or conditionally committed by another funder: none. Both invitations decide after QBE's 13 November pre-condition date.")
L.append("## 3. Buyers and grants received, all time (Goods-tagged)")
def block(start_label,stop_labels):
    out=[]; on=False
    for r in range(1,A.max_row+1):
        v=A.cell(r,1).value
        if v==start_label: on=True; continue
        if on and (v in stop_labels or (isinstance(v,str) and v.startswith("2. "))): break
        if on and v and A.cell(r,1).font.bold is False and A.cell(r,5).value: out.append([v,m(A.cell(r,2).value),A.cell(r,3).value,A.cell(r,4).value,A.cell(r,5).value])
    return out
buyers=block("Who",["Grant received (Goods-tagged)"]); grants=block("Grant received (Goods-tagged)",["To confirm with Ben (paid to the ACT ledger, not Goods-tagged in the mirror)"]); uncl=block("To confirm with Ben (paid to the ACT ledger, not Goods-tagged in the mirror)",[])
L.append("### Buyers (beds and washing machines)")
L.append(table(["Who","Gross AUD","Paper","What / status","Lane"],buyers))
tot=sum(float(A.cell(r,2).value) for r in range(1,A.max_row+1) if A.cell(r,5).value=="earned" and isinstance(A.cell(r,2).value,(int,float)))
L.append(f"Buyers paid, total: {m(tot)}. Beds sold and paid to date: 320 across four organisations, $197,060 ex GST on the bed lines alone.")
L.append("Demand not yet on paper: Centrecorp 130 Stretch Beds on quote (QU-0014); more than 200 requests each in Tennant Creek and Mparntwe; ALIVE's four community visits. None of these is an order.")
L.append("### Grants received")
L.append(table(["Funder","Gross AUD","Paper","What","Lane"],grants))
L.append("### Receipts still to classify (paid to the ACT ledger, not Goods-tagged)")
L.append(table(["Who","Gross AUD","Paper","Question","Lane"],uncl))
L.append("## 4. Outgoings: what the programme costs")
L.append(table(["Line","Amount","Funder today"],[
 ["Make and freight 400 first-stock beds at Witta (NM route)",m(P["B6"].value),"Paid for at $750 a bed by catalytic capital: TFFF 133, BMD 80, 187 to find"],
 ["Facilitation, four bed communities at $10,000",m(P["B7"].value),"BMD $40,000"],
 ["Two community plants (Palm Island, Maningrida, may change)",m(P["B8"].value),"QBE $300,000"],
 ["Witta rent and upkeep (no coordinator, no trainer: labour is in the bed)",m(P["B9"].value),"What stays in every bed, plus Goods' other income"],
 ["Founders, travel, accounting, marketing $10,000 (no head office)",m(P["B10"].value),"What stays in every bed, TFFF operating support, Goods' other income"],
 ["All lines added together (reference only, not how the business is priced)",m(P["B11"].value),"The business runs on $750 a bed plus what else it sells"],
 ["Spend falling in FY27 (Dec 2026 to Jun 2027)",m(A["B66"].value if isinstance(A["B66"].value,(int,float)) else next((A.cell(r,2).value for r in range(1,A.max_row+1) if str(A.cell(r,1).value).startswith("Spend falling")),"")),"Month 1 = December 2026"]]))
L.append("## 5. Options")
rows=[]
for r in range(6,11):
    rows.append([C.cell(r,1).value,m(C.cell(r,2).value),m(C.cell(r,3).value),C.cell(r,8).value,m(C.cell(r,9).value),m(C.cell(r,10).value),m(C.cell(r,12).value)])
L.append(table(["Case","QBE","Loan","Year 1 beds","Goods spend","Year 1 receipts","Peak cash need"],rows))
L.append("Full = QBE builds both plants, TFFF, BMD and the Sefa loan all land. Smaller = one plant and a smaller loan. No QBE = the plants wait. Receipts include invitations and a loan that are not signed; peak cash need is the timing measure.")
L.append("## 6. Balances")
L.append(table(["Entity","Amount","As at","Note"],[
 ["Nic sole trader ledger (where Goods trades today), cash","$143,997","9 Sep 2026","Xero"],
 ["Same ledger, receivables","$281,049","9 Sep 2026","Includes Rotary $82,500 bad debt and ALIVE $66,000"],
 ["Same ledger, payables","$311,237","9 Sep 2026","Xero"],
 ["Same ledger, FY27 cash income to date","$316,047","1 Jul to 9 Sep 2026","All ACT activity; Goods share not split"],
 ["The Butterfly Movement Ltd (applicant), cash","$4,041","30 Jun 2026","Unaudited; audit due mid-September"],
 ["The Butterfly Movement Ltd, net assets","$8,407","30 Jun 2026",""]]))
L.append("The Goods money sits in Nic's ledger. It needs a documented transfer before it is Butterfly's opening balance.")
L.append("## 7. Communities")
L.append(table(["Place","Role","Funded by"],[
 ["Alice Springs (Mparntwe), Oonchiumpa","Plant, project manager, youth employment programme","DEWR REAL $1.695M and NIAA, through Oonchiumpa. Outside this raise."],
 ["Palm Island","QBE plant 1, working choice","QBE $150,000"],
 ["Maningrida","QBE plant 2, working choice","QBE $150,000"],
 ["Four bed communities from: Mount Isa and Lower Gulf, Kununurra, Ceduna / Port Augusta, Tennant Creek, Utopia homelands","100 beds each, Nic and Ben facilitating","BMD, TFFF, buyers, still to raise"]]))
L.append("No community sees a price or a promise until they have agreed to it. This list is places in conversation, not commitments.")
L.append("## 8. Key dates")
L.append("- Fri 25 Sep 12pm: QBE application and BMD application due\n- Tue 6 or Wed 7 Oct: QBE review meeting (booked 7 Oct 9:45)\n- Fri 9 Oct 5pm: TFFF application due\n- Fri 23 Oct: QBE conditional outcomes\n- Fri 13 Nov: QBE pre-conditions due\n- 19 Nov: BMD board\n- Late Nov: TFFF board\n- Mon 14 Sep: Butterfly board, audited FY26 accounts")
L.append("## 9. Next decisions")
for r in range(1,A.max_row+1):
    v=A.cell(r,1).value
    if isinstance(v,str) and v[:2] in ("1.","2.","3.","4.","5.","6.") and A.cell(r,1).font.bold is False and len(v)>20: L.append("- "+v)
print("\n\n".join(L))
